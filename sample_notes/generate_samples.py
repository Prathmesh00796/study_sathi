"""
Script to generate sample DOCX and PDF files for StudySathi testing.
"""

from pathlib import Path
import docx

def create_sample_docx():
    notes_dir = Path(__file__).resolve().parent
    docx_path = notes_dir / "machine_learning_foundations.docx"

    doc = docx.Document()
    doc.add_heading("Machine Learning Foundations & Algorithms", level=0)

    p1 = doc.add_paragraph(
        "Supervised learning involves training a model on input features paired with known ground-truth labels. "
        "Common algorithms include Linear Regression for continuous variables and Logistic Regression for classification tasks."
    )

    doc.add_heading("1. Bias-Variance Tradeoff", level=1)
    doc.add_paragraph(
        "Bias refers to errors introduced by approximating a real-world problem with an overly simplistic model (underfitting). "
        "Variance refers to the model's excessive sensitivity to small fluctuations in the training set (overfitting). "
        "The optimal model minimizes total generalization error by balancing bias and variance."
    )

    doc.add_heading("2. Gradient Descent Optimization", level=1)
    doc.add_paragraph(
        "Gradient descent iteratively updates model parameters in the direction of steepest decrease of the cost function. "
        "Stochastic Gradient Descent (SGD) computes updates on single samples, Mini-Batch SGD operates on small subsets, "
        "and adaptive optimizers like Adam combine momentum with root mean square propagation."
    )

    doc.add_heading("3. Evaluation Metrics Table", level=1)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Metric"
    hdr_cells[1].text = "Formula"
    hdr_cells[2].text = "Use Case"

    data = [
        ("Precision", "TP / (TP + FP)", "Minimizing false positives (e.g. spam detection)"),
        ("Recall", "TP / (TP + FN)", "Minimizing false negatives (e.g. medical diagnosis)"),
        ("F1-Score", "2 * (P * R) / (P + R)", "Harmonic mean for imbalanced classification datasets"),
        ("ROC-AUC", "Area under TPR vs FPR curve", "Threshold-independent model discrimination ability")
    ]

    for metric, formula, use_case in data:
        row_cells = table.add_row().cells
        row_cells[0].text = metric
        row_cells[1].text = formula
        row_cells[2].text = use_case

    doc.save(str(docx_path))
    print(f"Generated {docx_path}")

def create_sample_pdf():
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    notes_dir = Path(__file__).resolve().parent
    pdf_path = notes_dir / "deep_learning_and_transformers.pdf"

    doc = SimpleDocTemplate(str(pdf_path), pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    title_style = styles["Title"]
    h1_style = styles["Heading1"]
    body_style = styles["Normal"]

    story.append(Paragraph("Deep Learning & Transformers Primer", title_style))
    story.append(Spacer(1, 14))

    story.append(Paragraph("1. Attention Is All You Need", h1_style))
    story.append(Paragraph(
        "The Transformer architecture introduced in 2017 discarded recurrent connections in favor of self-attention mechanisms. "
        "Self-attention allows tokens to compute relational weights against all other tokens in a sequence simultaneously in O(N^2) complexity.",
        body_style
    ))
    story.append(Spacer(1, 10))

    story.append(Paragraph("2. Multi-Head Attention Mechanics", h1_style))
    story.append(Paragraph(
        "Queries (Q), Keys (K), and Values (V) are projected into multiple representation subspaces. "
        "The scaled dot-product attention formula is: Attention(Q, K, V) = softmax( (Q * K^T) / sqrt(d_k) ) * V. "
        "Scaling by the square root of dimension d_k prevents softmax gradients from vanishing in regions with extremely large values.",
        body_style
    ))
    story.append(Spacer(1, 10))

    story.append(Paragraph("3. Pretraining and Fine-Tuning Paradigms", h1_style))
    story.append(Paragraph(
        "Modern LLMs undergo self-supervised pretraining on massive web-scale corpora, followed by Supervised Fine-Tuning (SFT) "
        "and Reinforcement Learning from Human Feedback (RLHF) or Direct Preference Optimization (DPO) to align model outputs with human intent.",
        body_style
    ))

    doc.build(story)
    print(f"Generated {pdf_path}")

if __name__ == "__main__":
    create_sample_docx()
    create_sample_pdf()


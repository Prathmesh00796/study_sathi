"""
Document Loader module for StudySathi.
Extracts and chunks text from PDF, DOCX, and TXT files with metadata preservation.
"""

import io
from pathlib import Path
from typing import List, Dict, Any, Union, Optional
from pypdf import PdfReader
import docx
from langchain_text_splitters import RecursiveCharacterTextSplitter
from config import CHUNK_SIZE, CHUNK_OVERLAP


# Extract raw text and page-level metadata from a PDF file
def extract_text_from_pdf(file_source: Union[str, Path, io.BytesIO, bytes], filename: str = "document.pdf") -> List[Dict[str, Any]]:
    """Reads a PDF and returns a list of page objects with text and page numbers."""
    pages_data = []
    try:
        if isinstance(file_source, bytes):
            stream = io.BytesIO(file_source)
        elif isinstance(file_source, (str, Path)):
            stream = open(file_source, "rb")
        else:
            stream = file_source

        reader = PdfReader(stream)
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                pages_data.append({
                    "text": text,
                    "metadata": {
                        "source": filename,
                        "page": idx + 1,
                        "file_type": "pdf",
                    }
                })

        if isinstance(file_source, (str, Path)) and hasattr(stream, "close"):
            stream.close()
    except Exception as e:
        pages_data.append({
            "text": f"[Error reading PDF {filename}: {str(e)}]",
            "metadata": {"source": filename, "page": 1, "error": str(e), "file_type": "pdf"}
        })
    return pages_data


# Extract raw text from a Microsoft Word (.docx) file
def extract_text_from_docx(file_source: Union[str, Path, io.BytesIO, bytes], filename: str = "document.docx") -> List[Dict[str, Any]]:
    """Reads a DOCX document including paragraphs and table contents."""
    try:
        if isinstance(file_source, bytes):
            stream = io.BytesIO(file_source)
        elif isinstance(file_source, (str, Path)):
            stream = open(file_source, "rb")
        else:
            stream = file_source

        doc = docx.Document(stream)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)
        if isinstance(file_source, (str, Path)) and hasattr(stream, "close"):
            stream.close()

        if not full_text.strip():
            return []

        return [{
            "text": full_text,
            "metadata": {
                "source": filename,
                "page": 1,
                "file_type": "docx",
            }
        }]
    except Exception as e:
        return [{
            "text": f"[Error reading DOCX {filename}: {str(e)}]",
            "metadata": {"source": filename, "page": 1, "error": str(e), "file_type": "docx"}
        }]


# Extract raw text from a plain text or Markdown file
def extract_text_from_txt(file_source: Union[str, Path, io.BytesIO, bytes], filename: str = "document.txt") -> List[Dict[str, Any]]:
    """Reads a TXT file with UTF-8 encoding (and latin-1 fallback)."""
    try:
        if isinstance(file_source, bytes):
            content = file_source
        elif isinstance(file_source, (str, Path)):
            with open(file_source, "rb") as f:
                content = f.read()
        else:
            content = file_source.read()

        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            text = content.decode("latin-1", errors="replace")

        if not text.strip():
            return []

        return [{
            "text": text.strip(),
            "metadata": {
                "source": filename,
                "page": 1,
                "file_type": "txt",
            }
        }]
    except Exception as e:
        return [{
            "text": f"[Error reading TXT {filename}: {str(e)}]",
            "metadata": {"source": filename, "page": 1, "error": str(e), "file_type": "txt"}
        }]


# Dispatch file parsing based on extension
def load_document(file_source: Union[str, Path, Any], filename: Optional[str] = None) -> List[Dict[str, Any]]:
    """Loads a single document and returns list of raw pages/sections."""
    if filename is None:
        if hasattr(file_source, "name"):
            filename = file_source.name
        elif isinstance(file_source, (str, Path)):
            filename = Path(file_source).name
        else:
            filename = "uploaded_file.txt"

    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_source, filename=filename)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(file_source, filename=filename)
    elif ext in (".txt", ".md", ".csv", ".json"):
        return extract_text_from_txt(file_source, filename=filename)
    else:
        # Default fallback to plain text reader
        return extract_text_from_txt(file_source, filename=filename)


# Split extracted document pages into semantic chunks for vector indexing
def chunk_documents(
    documents: List[Dict[str, Any]],
    chunk_size: int = CHUNK_SIZE,
    chunk_overlap: int = CHUNK_OVERLAP,
) -> List[Dict[str, Any]]:
    """Splits document text into chunks while preserving and enhancing metadata."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    chunks = []
    chunk_counter = 0

    for doc in documents:
        text = doc.get("text", "")
        base_meta = doc.get("metadata", {})
        split_texts = splitter.split_text(text)

        for i, chunk_text in enumerate(split_texts):
            clean_chunk = chunk_text.strip()
            if not clean_chunk:
                continue

            chunk_meta = dict(base_meta)
            chunk_meta["chunk_index"] = i
            chunk_meta["chunk_id"] = f"{base_meta.get('source', 'doc')}_p{base_meta.get('page', 1)}_c{i}"
            chunk_meta["char_count"] = len(clean_chunk)

            chunks.append({
                "id": f"chunk_{chunk_counter}",
                "text": clean_chunk,
                "metadata": chunk_meta,
            })
            chunk_counter += 1

    return chunks
